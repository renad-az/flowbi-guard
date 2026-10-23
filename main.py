from pathlib import Path
from openpyxl import load_workbook

project_folder = Path(__file__).resolve().parent

source_file = (
    project_folder
    / "data"
    / "incoming"
    / "FlowBI_Project_Management.xlsx"
)

workbook = load_workbook(source_file, read_only=True)

print("Source:", source_file)
print("Sheets:", workbook.sheetnames)

workbook.close()